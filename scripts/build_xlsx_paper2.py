# -*- coding: utf-8 -*-
"""بناء نسخة Excel من بحث «خصائص لغة الإعلام» + ورقة ملاحظات المراجعة.
يستخرج النص مباشرة من HTML النهائي لضمان التطابق الحرفي."""
import sys, os, re, html as ihtml, math

XLSX_SKILL_DIR = "/home/z/my-project/skills/xlsx"
for sub in [XLSX_SKILL_DIR, os.path.join(XLSX_SKILL_DIR, "templates")]:
    if sub not in sys.path:
        sys.path.insert(0, sub)

import base
# خط عربي آمن على كل الأنظمة (Arial يدعم العربية كلياً) — عبر متغير المكتبة الواحد
base.FONT_NAME = "Arial"
base.HEADER_BOLD = True

from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

HTML_PATH = "/home/z/my-project/download/بحث_خصائص_لغة_الإعلام.html"
OUT_PATH = "/home/z/my-project/download/بحث_خصائص_لغة_الإعلام.xlsx"

# ---------------------------------------------------------------- استخراج النص
src = open(HTML_PATH, encoding="utf-8").read()

def clean(s):
    s = re.sub(r"<br\s*/?>", " ", s)
    s = re.sub(r"<[^>]+>", "", s)
    s = ihtml.unescape(s)
    return re.sub(r"\s+", " ", s).strip()

# 1) الغلاف: بيانات المعهد + أزواج label/value
inst = clean(re.search(r'<div class="inst-box">(.*?)</div>', src, re.S).group(1))
cover_pairs = [
    (clean(l), clean(v))
    for l, v in re.findall(
        r'<div class="cf-label">(.*?)</div>\s*<div class="cf-value[^"]*">(.*?)</div>', src, re.S)
]

# 2) الفهرس (سطر واحد مدمج)
toc_items = [
    (clean(t), clean(p))
    for t, p in re.findall(
        r'<div class="toc-row">.*?<div class="toc-text">(.*?)</div>.*?<div class="toc-page">(.*?)</div>',
        src, re.S)
]
toc_line = " | ".join(f"{t} {p}" for t, p in toc_items)

# 3) أقسام المتن: h2/h3/p بالترتيب داخل كل section
PAGE_MAP = {"المقدمة": "3", "المبحث الأول": "4", "المبحث الثاني": "5",
            "المبحث الثالث": "6", "الخاتمة": "7", "قائمة المراجع": "7"}
sections = re.findall(r'<section class="page[^"]*">(.*?)</section>', src, re.S)
rows_body = []
for sec in sections:
    if 'class="toc-title"' in sec:
        continue  # صفحة الفهرس مغطاة بسطر الفهرس المدمج
    cur_h2, cur_h3 = None, None
    for m in re.finditer(r'<h2[^>]*>(.*?)</h2>|<h3[^>]*>(.*?)</h3>|<p[^>]*>(.*?)</p>', sec, re.S):
        h2, h3, p = m.groups()
        if h2 is not None:
            cur_h2, cur_h3 = clean(h2).rstrip(":"), None
        elif h3 is not None:
            cur_h3 = clean(h3).rstrip(":")
        else:
            text = clean(p)
            if not text:
                continue
            label = cur_h3 if cur_h3 else (cur_h2 or "")
            page = next((v for k, v in PAGE_MAP.items() if cur_h2 and cur_h2.startswith(k)), "")
            rows_body.append((page, label, text))

rows = [("الغلاف", "بيانات المعهد", inst)]
rows += [("الغلاف", l.rstrip(" :"), v) for l, v in cover_pairs]
rows += [("2", "الفهرس", toc_line)]
rows += rows_body

# --------------------------------------------- فحص ذاتي: البنية والعبارات الجديدة
assert len(rows) == 33, f"عدد الصفوف غير متوقع: {len(rows)}"
joined = " ".join(t for _, _, t in rows)
for phrase in ["من غير اللغة لا يستطيع", "حسن عماد مكاوي في كتابه", "أهم سلاح",
               "لماذا نشعر أحياناً", "وكالات الأنباء", "وخلاصة رأيي", "التساؤلات التالية"]:
    assert phrase in joined, f"عبارة مفقودة من النسخة المحدثة: {phrase}"
assert "تعد لغة الإعلام من المواضيع" not in joined, "بقيت الصياغة القديمة!"
print(f"Self-check OK: {len(rows)} rows, all new phrases present")

# ---------------------------------------------------------------- بناء الملف
wb = Workbook()

# ===== الورقة 1: البحث =====
ws = wb.active
ws.title = "البحث"
ws.sheet_view.rightToLeft = True
HDRS1, W1 = ["الصفحة", "القسم", "النص"], [9, 26, 96]
last_col1 = 1 + len(HDRS1)  # D
base.setup_sheet(ws, title="بحث خصائص لغة الإعلام — نص كامل قابل للنسخ والتعديل", last_col=last_col1)
ws.cell(row=2, column=2).alignment = Alignment(horizontal="right", vertical="center")
for i, h in enumerate(HDRS1, 2):
    ws.cell(row=4, column=i, value=h)
base.style_header_row(ws, 4, 2, last_col1)

def est_height(pairs, base_h=22):
    """تقدير ارتفاع الصف من عدد الأسطر الملغومة في كل عمود نصي."""
    lines = 1
    for text, width in pairs:
        cpl = max(20, int(width * 0.82))
        lines = max(lines, math.ceil(len(text) / cpl))
    return max(base_h, lines * 15.5 + 8)

r = 5
for i, (page, label, text) in enumerate(rows):
    ws.cell(row=r, column=2, value=page)
    ws.cell(row=r, column=3, value=label)
    ws.cell(row=r, column=4, value=text)
    base.style_data_row(ws, r, 2, last_col1, i)
    ws.cell(row=r, column=2).alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(row=r, column=3).alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
    ws.cell(row=r, column=4).alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
    ws.row_dimensions[r].height = est_height([(label, W1[1]), (text, W1[2])])
    r += 1
for col, w in zip("BCD", W1):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "B5"
cap = ws.cell(row=r + 1, column=2, value="نسخة نصية من البحث للنسخ والتعديل السهل — النسخة الرسمية للتسليم هي ملف PDF بنفس الاسم.")
cap.font = base.font_caption()
cap.alignment = Alignment(horizontal="right", vertical="center")

# ===== الورقة 2: ملاحظات المراجعة =====
ws2 = wb.create_sheet("ملاحظات المراجعة")
ws2.sheet_view.rightToLeft = True
HDRS2, W2 = ["الموقع", "الملاحظة من الفاحص", "التصنيف", "الإجراء المتخذ"], [18, 46, 10, 52]
last_col2 = 1 + len(HDRS2)  # E
base.setup_sheet(ws2, title="ملاحظات مراجعة البحث (تقييم الفاحص 7.5/10) والإجراءات", last_col=last_col2)
ws2.cell(row=2, column=2).alignment = Alignment(horizontal="right", vertical="center")
for i, h in enumerate(HDRS2, 2):
    ws2.cell(row=4, column=i, value=h)
base.style_header_row(ws2, 4, 2, last_col2)

notes = [
    ("المقدمة (ص3)", "الصياغة الافتتاحية «تعد لغة الإعلام من المواضيع المهمة...» عامة ورسمية، وبقية العناصر (أهمية/أهداف/منهج) بأسلوب مرتب ومتوقع.", "أساسي",
     "أعيدت الافتتاحية لتبدأ مباشرة: «من غير اللغة لا يستطيع الإعلام أن يوصل كلامه إلى الناس...» مع ربط الجمل في تدفق واحد أقل تقسيماً، وتبسيط صيغة مشكلة البحث."),
    ("المبحث الأول (ص4)", "التعريف والحديث عن الفصحى والعامية مكتوبان بطريقة تلخيصية عامة، من غير أمثلة أو إحالات لمصادر داخل النص.", "أساسي",
     "أضيف مثال ملموس في التعريف (صياغة الخبر الواحد بكلمات قريبة من كلام الناس) وإحالة داخل النص إلى كتاب حسن عماد مكاوي «لغة الإعلام» الموجود في قائمة المراجع."),
    ("المبحث الثاني (ص5)", "كل خاصية معروضة بنفس النمط تقريباً: تعريف مختصر ثم تفسير ثم نتيجة.", "أساسي",
     "كسر النمط: بداية جديدة لخاصية التشويق («والتشويق أهم سلاح عند المذيع أو الكاتب»)، وتحويل خاصية القرب من الحياة إلى صيغة سؤال وإجابة."),
    ("المبحث الثالث (ص6)", "التوصيات عامة ومختصرة، ومفيش أمثلة عملية بتدعمها.", "أساسي",
     "أضيف مثالاً عملياً: المذيع لا يقرأ الخبر كما يأتي من وكالات الأنباء بل يعيد صياغته بكلمات أسهل يفهمها الجمهور."),
    ("الخاتمة (ص7)", "الخاتمة تعيد معظم النقاط المذكورة من غير تحليل أو استنتاج شخصي واضح.", "أساسي",
     "أضيف استنتاجاً شخصياً في نهاية النتائج: «وخلاصة رأيي أن هذه اللغة حية تتبع الناس، فكل ما تغير في حياتهم وكلامهم انعكس على لغة الإعلام»."),
    ("قائمة المراجع (ص7)", "تأكد أن كل مرجع موجود فعلاً وأن بيانات الكتاب والناشر وسنة النشر صحيحة.", "ثانوي",
     "المراجع كتب حقيقية معروفة في أدبيات الإعلام العربي (عبد المطلب عبد حميد، إبراهيم إمام، سعيد سلامة، حسن عماد مكاوي وغيرهم). يُنصح بمطابقة سنة النشر والناشر على النسخ الورقية المتاحة في مكتبة المعهد قبل التسليم."),
    ("التحليل العام", "البحث يعتمد على الشرح العام أكتر من الأمثلة والتحليل.", "ثانوي",
     "أضيفت أمثلة محسوسة في ص4 وص6 مع إبقاء البنية الأكاديمية المطلوبة لمبحث/خاتمة بحث جامعي."),
    ("اللغة", "تكرار تعبيرات مثل «ولذلك» و«ولهذا» و«ومن أبرز» يجعل الأسلوب متوقعاً.", "ثانوي",
     "تنويع الروابط: حذف «ولهذا» من المتن، وبقاء «ولذلك» مرة واحدة فقط، واستبدال «ومن أبرز» بصيغ متنوعة، وإزالة تكرار كلمة «أبرز» من فقرات المتن."),
]
r = 5
for i, (loc, note, kind, action) in enumerate(notes):
    ws2.cell(row=r, column=2, value=loc)
    ws2.cell(row=r, column=3, value=note)
    ws2.cell(row=r, column=4, value=kind)
    ws2.cell(row=r, column=5, value=action)
    base.style_data_row(ws2, r, 2, last_col2, i)
    for c in (2, 3, 4, 5):
        ws2.cell(row=r, column=c).alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
    ws2.cell(row=r, column=4).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws2.row_dimensions[r].height = est_height([(note, W2[1]), (action, W2[3])])
    r += 1
for col, w in zip("BCDE", W2):
    ws2.column_dimensions[col].width = w
ws2.freeze_panes = "B5"
cap2 = ws2.cell(row=r + 1, column=2, value="المصدر: تقرير الفاحص الخارجي المرسل من الطالب + التعديلات المنفذة فعلياً على نسخة HTML/PDF.")
cap2.font = base.font_caption()
cap2.alignment = Alignment(horizontal="right", vertical="center")

wb.properties.creator = "Z.ai"
wb.save(OUT_PATH)
print("Saved:", OUT_PATH)
